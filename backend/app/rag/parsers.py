from dataclasses import dataclass
from io import BytesIO
from time import monotonic

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import settings


@dataclass(frozen=True)
class ParsedBlock:
    text: str
    page: int | None = None
    paragraph: int | None = None
    block_index: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    blocks: list[ParsedBlock]


def parse_document_bytes(*, content: bytes, mime: str | None, filename: str) -> str:
    return parse_document_blocks(content=content, mime=mime, filename=filename).text


def parse_document_blocks(
    *, content: bytes, mime: str | None, filename: str
) -> ParsedDocument:
    normalized_mime = (mime or "").lower()
    normalized_name = filename.lower()

    if normalized_mime in {"text/plain", "text/markdown"} or normalized_name.endswith(
        (".txt", ".md")
    ):
        return parse_text_blocks(content)
    if normalized_mime == "application/pdf" or normalized_name.endswith(".pdf"):
        return parse_pdf_blocks(content)
    if (
        normalized_mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or normalized_name.endswith(".docx")
    ):
        return parse_docx_blocks(content)

    if normalized_mime.startswith("text/"):
        return parse_text_blocks(content)
    raise ValueError("unsupported_document_type")


def parse_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def parse_text_blocks(content: bytes) -> ParsedDocument:
    text = parse_text(content)
    paragraphs = split_text_blocks(text)
    blocks = [
        ParsedBlock(text=paragraph, paragraph=index, block_index=index)
        for index, paragraph in enumerate(paragraphs, start=1)
    ]
    return ParsedDocument(text="\n\n".join(paragraphs), blocks=blocks)


def parse_pdf(content: bytes) -> str:
    return parse_pdf_blocks(content).text


def parse_pdf_blocks(content: bytes) -> ParsedDocument:
    deadline = monotonic() + settings.parser_timeout_seconds
    reader = PdfReader(BytesIO(content))
    if len(reader.pages) > settings.parser_pdf_max_pages:
        raise ValueError("parser_failed")
    blocks: list[ParsedBlock] = []
    paragraph_index = 0
    block_index = 0
    for page_index, page in enumerate(reader.pages, start=1):
        ensure_parse_deadline(deadline)
        text = page.extract_text() or ""
        for paragraph in split_text_blocks(text):
            paragraph_index += 1
            block_index += 1
            blocks.append(
                ParsedBlock(
                    text=paragraph,
                    page=page_index,
                    paragraph=paragraph_index,
                    block_index=block_index,
                )
            )
    return ParsedDocument(text="\n\n".join(block.text for block in blocks), blocks=blocks)


def parse_docx(content: bytes) -> str:
    return parse_docx_blocks(content).text


def parse_docx_blocks(content: bytes) -> ParsedDocument:
    deadline = monotonic() + settings.parser_timeout_seconds
    document = DocxDocument(BytesIO(content))
    blocks: list[ParsedBlock] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        ensure_parse_deadline(deadline)
        text = paragraph.text.strip()
        if text:
            block_index = len(blocks) + 1
            blocks.append(
                ParsedBlock(text=text, paragraph=block_index, block_index=block_index)
            )
    return ParsedDocument(text="\n\n".join(block.text for block in blocks), blocks=blocks)


def split_text_blocks(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [
        "\n".join(line.strip() for line in paragraph.splitlines() if line.strip())
        for paragraph in normalized.split("\n\n")
    ]
    blocks = [paragraph for paragraph in paragraphs if paragraph.strip()]
    if blocks:
        return blocks
    return [line.strip() for line in normalized.splitlines() if line.strip()]


def ensure_parse_deadline(deadline: float) -> None:
    if monotonic() > deadline:
        raise ValueError("parser_failed")
