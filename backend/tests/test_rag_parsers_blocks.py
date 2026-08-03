from io import BytesIO
from types import SimpleNamespace

from docx import Document as DocxDocument

from app.rag.parsers import parse_document_blocks, parse_document_bytes


def test_parse_txt_blocks_assigns_paragraph_and_block_index() -> None:
    parsed = parse_document_blocks(
        content=b"alpha\n\nbeta\n\n\ngamma",
        mime="text/plain",
        filename="sample.txt",
    )

    assert parsed.text == "alpha\n\nbeta\n\ngamma"
    assert [block.text for block in parsed.blocks] == ["alpha", "beta", "gamma"]
    assert [block.paragraph for block in parsed.blocks] == [1, 2, 3]
    assert [block.block_index for block in parsed.blocks] == [1, 2, 3]
    assert all(block.page is None for block in parsed.blocks)


def test_parse_markdown_blocks_keeps_markdown_text() -> None:
    parsed = parse_document_blocks(
        content="# Root\nIntro\n\n## Child\nDetails".encode(),
        mime="text/markdown",
        filename="sample.md",
    )

    assert len(parsed.blocks) == 2
    assert parsed.blocks[0].text == "# Root\nIntro"
    assert parsed.blocks[1].paragraph == 2
    assert parse_document_bytes(
        content=b"# Root\nIntro", mime="text/markdown", filename="sample.md"
    ) == "# Root\nIntro"


def test_parse_docx_blocks_assigns_paragraphs() -> None:
    document = DocxDocument()
    document.add_paragraph("Alpha")
    document.add_paragraph("")
    document.add_paragraph("Beta")
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document_blocks(
        content=buffer.getvalue(),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="sample.docx",
    )

    assert [block.text for block in parsed.blocks] == ["Alpha", "Beta"]
    assert [block.paragraph for block in parsed.blocks] == [1, 2]
    assert [block.block_index for block in parsed.blocks] == [1, 2]


def test_parse_pdf_blocks_assigns_page(monkeypatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakePdfReader:
        def __init__(self, stream) -> None:
            self.pages = [FakePage("Page one"), FakePage("Page two\n\nMore")]

    monkeypatch.setattr("app.rag.parsers.PdfReader", FakePdfReader)
    monkeypatch.setattr(
        "app.rag.parsers.settings",
        SimpleNamespace(parser_timeout_seconds=30, parser_pdf_max_pages=10),
    )

    parsed = parse_document_blocks(
        content=b"%PDF",
        mime="application/pdf",
        filename="sample.pdf",
    )

    assert [block.page for block in parsed.blocks] == [1, 2, 2]
    assert [block.block_index for block in parsed.blocks] == [1, 2, 3]
    assert parsed.text == "Page one\n\nPage two\n\nMore"
