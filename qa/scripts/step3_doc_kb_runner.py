import json
import os
import subprocess
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "qa"
DOCS = QA / "scripts" / "fixtures_docs"
ENV = QA / "env"
RECORDS = QA / "records"
BASE = "http://localhost:8001/api/v1"


def parse_json(text: str):
    try:
        return json.loads(text)
    except Exception:
        return None


def req(method: str, path: str, *, token=None, json_body=None, raw_body=None, content_type=None, timeout=90):
    url = path if path.startswith("http") else f"{BASE}{path}"
    headers = {}
    data = None
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if json_body is not None:
        data = json.dumps(json_body).encode()
        headers["Content-Type"] = "application/json"
    if raw_body is not None:
        data = raw_body
        headers["Content-Type"] = content_type
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as resp:
            text = resp.read().decode("utf-8", errors="replace")
            return {"status": resp.status, "body": parse_json(text) if parse_json(text) is not None else text}
    except urllib.error.HTTPError as exc:
        text = exc.read().decode("utf-8", errors="replace")
        return {"status": exc.code, "body": parse_json(text) if parse_json(text) is not None else text}


def detail(resp):
    b = resp["body"]
    return b.get("detail") if isinstance(b, dict) else b


def run(cmd, *, input_text=None, timeout=120):
    p = subprocess.run(cmd, cwd=ROOT, input=input_text, text=True, capture_output=True, timeout=timeout)
    return {"code": p.returncode, "stdout": p.stdout, "stderr": p.stderr}


def compose_exec(service, args, *, input_text=None, timeout=120):
    return run(["docker", "compose", "exec", "-T", service, *args], input_text=input_text, timeout=timeout)


def psql(sql):
    res = compose_exec("postgres", ["psql", "-U", "app", "-d", "eap", "-A", "-F", "\t", "-q", "-c", sql])
    lines = [line for line in res["stdout"].splitlines() if line and not line.startswith("(")]
    if not lines:
        return []
    header = lines[0].split("\t")
    out = []
    for line in lines[1:]:
        cols = line.split("\t")
        out.append(dict(zip(header, cols)))
    return out


def make_docs():
    DOCS.mkdir(parents=True, exist_ok=True)
    normal_text = (
        "Enterprise contract approval checklist.\n\n"
        "Review legal entity, authorization, payment terms, delivery acceptance, confidentiality, data security, breach liability and dispute resolution.\n\n"
        "Approvers should verify budget source, seal usage, attachments and timeline owners."
    )
    (DOCS / "normal.txt").write_text(normal_text, encoding="utf-8")
    (DOCS / "normal.md").write_text("# Contract Review\n\n" + normal_text + "\n\n- item one\n- item two\n", encoding="utf-8")
    (DOCS / "empty.txt").write_bytes(b"")
    long_para = ("0123456789 abcdefghijklmnopqrstuvwxyz contract approval risk control data security. " * 2600)
    (DOCS / "long.txt").write_text(long_para, encoding="utf-8")
    (DOCS / "broken.pdf").write_bytes(b"%PDF-broken\nnot a valid pdf")
    docx_script = f"""
import base64
from io import BytesIO
from docx import Document
d=Document()
d.add_heading('QA DOCX Contract Review', 1)
d.add_paragraph({normal_text!r})
d.add_paragraph('Second paragraph for chunk extraction and embedding.')
buf=BytesIO()
d.save(buf)
print(base64.b64encode(buf.getvalue()).decode())
"""
    docx_res = compose_exec("backend", ["python", "-"], input_text=docx_script)
    import base64
    (DOCS / "normal.docx").write_bytes(base64.b64decode(docx_res["stdout"].strip()))
    make_minimal_pdf(DOCS / "normal.pdf", "QA PDF Contract Review\\nPayment terms and acceptance checks.\\nData security and dispute resolution.")


def make_minimal_pdf(path: Path, text: str):
    # Minimal single-page PDF with simple text operators. pypdf extracts this text.
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").replace("\n", ") Tj T* (")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin1")
    objects = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{i} 0 obj\n".encode())
        content.extend(obj)
        content.extend(b"\nendobj\n")
    xref = len(content)
    content.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for off in offsets[1:]:
        content.extend(f"{off:010d} 00000 n \n".encode())
    content.extend(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(bytes(content))


def login_fixture():
    fx = json.loads((ENV / "step1-fixtures.json").read_text())
    alice = fx["users"]["alice"]
    tenant = fx["tenants"]["A"]
    resp = req("POST", "/auth/login", json_body={"tenant_code": tenant["code"], "username": alice["username"], "password": alice["password"]})
    if resp["status"] != 200:
        raise RuntimeError(f"alice login failed: {resp}")
    return resp["body"]["access_token"], fx


def upload(token, kb_id, file_path: Path):
    boundary = f"----qa{int(time.time()*1000)}"
    data = file_path.read_bytes()
    mime = guess_mime(file_path)
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'
        f"Content-Type: {mime}\r\n\r\n"
    ).encode() + data + f"\r\n--{boundary}--\r\n".encode()
    return req("POST", f"/kbs/{kb_id}/documents", token=token, raw_body=body, content_type=f"multipart/form-data; boundary={boundary}", timeout=120)


def guess_mime(path: Path):
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pdf": "application/pdf",
    }.get(path.suffix.lower(), "application/octet-stream")


def create_kb(token, name, *, config=None, embedding_model="mock-embedding"):
    return req("POST", "/kbs", token=token, json_body={"name": name, "type": "faq", "description": "qa-step3", "config": config or {}, "embedding_model": embedding_model}, timeout=120)


def wait_doc(token, doc_id, timeout=90):
    start = time.time()
    last = None
    while time.time() - start < timeout:
        rows = psql(f"select parse_status, meta::text from documents where id='{doc_id}';")
        if rows:
            last = rows[0]
            if last["parse_status"] in {"done", "failed"}:
                return last
        time.sleep(0.5)
    return last or {"parse_status": "timeout", "meta": "{}"}


def chunk_counts(doc_id):
    out = {}
    for table in ["chunks", "chunks_1024", "chunks_3072"]:
        rows = psql(f"select count(*) as count, coalesce(max(vector_dims(embedding)),0) as dim, coalesce(min(seq),-1) as min_seq, coalesce(max(seq),-1) as max_seq, coalesce(min(length(content)),0) as min_len, coalesce(max(length(content)),0) as max_len from {table} where doc_id='{doc_id}';")
        out[table] = rows[0] if rows else {}
    return out


def doc_row(doc_id):
    rows = psql(f"select id::text,kb_id::text,source_uri,parse_status,meta::text from documents where id='{doc_id}';")
    return rows[0] if rows else {}


def force_reindex(token, kb_id, doc_id):
    before = chunk_counts(doc_id)
    resp = req("POST", f"/kbs/{kb_id}/reindex", token=token, json_body={})
    status = wait_doc(token, doc_id, timeout=120)
    after = chunk_counts(doc_id)
    return {"response": resp, "before": before, "status": status, "after": after}


def delete_minio_source(document_id):
    row = doc_row(document_id)
    uri = row.get("source_uri") or ""
    script = f"""
from app.core.storage import delete_object
delete_object({uri!r})
print('deleted')
"""
    return compose_exec("backend", ["python", "-"], input_text=script)


def run_chunk_model_unit():
    script = """
from app.rag.chunk_store import chunk_model_for_dim
cases = [(1024, 'Chunk1024'), (1536, 'Chunk'), (3072, 'Chunk3072')]
for dim, name in cases:
    actual = chunk_model_for_dim(dim).__name__
    assert actual == name, (dim, actual, name)
try:
    chunk_model_for_dim(768)
except ValueError as exc:
    assert str(exc) == 'unsupported_embedding_dim:768', str(exc)
else:
    raise AssertionError('768 did not fail')
print('PASS')
"""
    return compose_exec("backend", ["python", "-"], input_text=script)


def main():
    make_docs()
    token, fixture = login_fixture()
    ts = str(int(time.time()))
    results = {"ts": ts, "maas_production_mode": compose_exec("maas", ["printenv", "PRODUCTION_MODE"])["stdout"].strip(), "groups": {}, "defects": []}

    # Group 1
    main_rows = []
    for filename in ["normal.txt", "normal.md", "normal.docx", "normal.pdf"]:
        kb = create_kb(token, f"QA Step3 main {filename} {ts}")
        up = upload(token, kb["body"]["id"], DOCS / filename)
        status = wait_doc(token, up["body"]["id"], timeout=120) if up["status"] == 201 else {"parse_status": "upload_failed", "meta": json.dumps(up)}
        counts = chunk_counts(up["body"]["id"]) if up["status"] == 201 else {}
        meta = parse_json(status.get("meta", "{}")) or {}
        count = int(counts.get("chunks", {}).get("count", "0") or 0)
        ok = up["status"] == 201 and status["parse_status"] == "done" and count > 0 and counts["chunks"]["dim"] == "1536" and counts["chunks_1024"]["count"] == "0" and counts["chunks_3072"]["count"] == "0"
        main_rows.append({"file": filename, "upload": up["status"], "parse_status": status["parse_status"], "meta": meta, "counts": counts, "result": "PASS" if ok else "FAIL"})
    results["groups"]["main"] = main_rows

    # Group 2
    kb_small = create_kb(token, f"QA Step3 chunk small {ts}", config={"chunk_size": 400, "overlap": 40})
    kb_large = create_kb(token, f"QA Step3 chunk large {ts}", config={"chunk_size": 1000, "overlap": 100})
    up_small = upload(token, kb_small["body"]["id"], DOCS / "long.txt")
    up_large = upload(token, kb_large["body"]["id"], DOCS / "long.txt")
    st_small = wait_doc(token, up_small["body"]["id"], 180)
    st_large = wait_doc(token, up_large["body"]["id"], 180)
    cc_small = chunk_counts(up_small["body"]["id"])
    cc_large = chunk_counts(up_large["body"]["id"])
    empty_kb = create_kb(token, f"QA Step3 empty {ts}")
    empty_up = upload(token, empty_kb["body"]["id"], DOCS / "empty.txt")
    empty_st = wait_doc(token, empty_up["body"]["id"], 90)
    empty_counts = chunk_counts(empty_up["body"]["id"])
    results["groups"]["chunking"] = {
        "small": {"status": st_small, "counts": cc_small},
        "large": {"status": st_large, "counts": cc_large},
        "empty": {"upload": empty_up["status"], "status": empty_st, "counts": empty_counts},
        "result": "PASS" if int(cc_small["chunks"]["count"]) > int(cc_large["chunks"]["count"]) > 0 and empty_st["parse_status"] == "failed" and "document_has_no_text" in empty_st["meta"] and empty_counts["chunks"]["count"] == "0" else "FAIL",
    }

    # Group 3
    broken_kb = create_kb(token, f"QA Step3 broken pdf {ts}")
    broken_up = upload(token, broken_kb["body"]["id"], DOCS / "broken.pdf")
    broken_st = wait_doc(token, broken_up["body"]["id"], 90)
    long_kb = create_kb(token, f"QA Step3 long robust {ts}", config={"chunk_size": 800, "overlap": 80})
    t0 = time.time()
    long_up = upload(token, long_kb["body"]["id"], DOCS / "long.txt")
    long_st = wait_doc(token, long_up["body"]["id"], 240)
    long_elapsed = round(time.time() - t0, 2)
    long_counts = chunk_counts(long_up["body"]["id"])
    reindex = force_reindex(token, long_kb["body"]["id"], long_up["body"]["id"])
    delete_res = delete_minio_source(long_up["body"]["id"])
    reindex_missing = force_reindex(token, long_kb["body"]["id"], long_up["body"]["id"])
    results["groups"]["robust"] = {
        "broken_pdf": {"upload": broken_up["status"], "status": broken_st, "counts": chunk_counts(broken_up["body"]["id"])},
        "long_doc": {"upload": long_up["status"], "status": long_st, "counts": long_counts, "elapsed_sec": long_elapsed},
        "force_reindex": reindex,
        "missing_source_reindex": {"delete_source": delete_res, "reindex": reindex_missing},
    }

    # Group 4
    unit = run_chunk_model_unit()
    bad_dim_model = f"qa-step3-bad-dim-{ts}"
    bad_ch = req("POST", "http://localhost:8100/admin/channels", json_body={"model": bad_dim_model, "model_type": "embedding", "provider": "mock", "tenant_id": fixture["tenants"]["A"]["id"], "base_url": "mock://local", "api_key": "x", "weight": 1, "status": "active"})
    # No public API to set provider_config request_defaults dimensions=768 via channel create/update.
    bad_kb = create_kb(token, f"QA Step3 bad dim {ts}", embedding_model=bad_dim_model)
    results["groups"]["dimensions"] = {
        "unit": unit,
        "bad_dim_construct": {"channel": bad_ch, "kb_create": bad_kb, "note": "demo channel create cannot set request_defaults dimensions=768; mock embedding resolves to 1536, so unsupported dim API path is not constructible via public API"},
        "pending_real_key": "BGE 1024 / text-embedding-3-large 3072 E2E pending real embedding key",
        "result": "PASS" if unit["code"] == 0 and "PASS" in unit["stdout"] else "FAIL",
    }

    # Group 5
    lock_kb = create_kb(token, f"QA Step3 lock {ts}")
    lock_up = upload(token, lock_kb["body"]["id"], DOCS / "normal.txt")
    lock_st = wait_doc(token, lock_up["body"]["id"], 90)
    patch_locked = req("PATCH", f"/kbs/{lock_kb['body']['id']}", token=token, json_body={"embedding_model": bad_dim_model})
    empty_lock_kb = create_kb(token, f"QA Step3 empty lock {ts}")
    patch_empty = req("PATCH", f"/kbs/{empty_lock_kb['body']['id']}", token=token, json_body={"embedding_model": bad_dim_model})
    results["groups"]["locking"] = {
        "locked_kb_doc_status": lock_st,
        "patch_locked": patch_locked,
        "patch_empty": patch_empty,
        "result": "PASS" if patch_locked["status"] == 409 and detail(patch_locked) == "kb_embedding_model_locked_has_documents" and patch_empty["status"] == 200 else "FAIL",
    }

    assess_defects(results)
    (ENV / "step3-results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    md = render(results)
    (RECORDS / "step3-doc-kb.md").write_text(md, encoding="utf-8")
    print(md)


def assess_defects(results):
    for row in results["groups"]["main"]:
        if row["result"] != "PASS":
            results["defects"].append({"group": "main", "name": row["file"], "severity": "P1", "evidence": row})
    if results["groups"]["chunking"]["result"] != "PASS":
        results["defects"].append({"group": "chunking", "name": "chunking/empty behavior mismatch", "severity": "P1", "evidence": results["groups"]["chunking"]})
    rob = results["groups"]["robust"]
    if rob["broken_pdf"]["status"]["parse_status"] != "failed":
        results["defects"].append({"group": "robust", "name": "broken pdf did not fail cleanly", "severity": "P1", "evidence": rob["broken_pdf"]})
    if rob["long_doc"]["status"]["parse_status"] != "done":
        results["defects"].append({"group": "robust", "name": "long document not ingested", "severity": "P1", "evidence": rob["long_doc"]})
    before = int(rob["force_reindex"]["before"]["chunks"]["count"])
    after = int(rob["force_reindex"]["after"]["chunks"]["count"])
    if before != after or after <= 0:
        results["defects"].append({"group": "robust", "name": "force reindex duplicate/loss", "severity": "P1", "evidence": rob["force_reindex"]})
    if rob["missing_source_reindex"]["reindex"]["status"]["parse_status"] != "done":
        results["defects"].append({"group": "robust", "name": "missing source reindex fallback failed", "severity": "P1", "evidence": rob["missing_source_reindex"]})
    if results["groups"]["locking"]["result"] != "PASS":
        results["defects"].append({"group": "locking", "name": "embedding model lock behavior mismatch", "severity": "P1", "evidence": results["groups"]["locking"]})


def short(obj, max_len=180):
    text = json.dumps(obj, ensure_ascii=False) if not isinstance(obj, str) else obj
    text = text.replace("|", "\\|").replace("\n", " ")
    return text if len(text) <= max_len else text[: max_len - 3] + "..."


def render(results):
    lines = ["# Step 3 文档上传 / 解析 / 切片 / Embedding / 入库测试", "", f"执行时间：{time.strftime('%Y-%m-%d %H:%M:%S %z')}", f"MaaS PRODUCTION_MODE：`{results['maas_production_mode']}`", ""]
    lines += ["## 第 1 组：正常主链路 1536", "", "| 文件 | parse_status | chunk_count | 落表 | 维度 | 结论 |", "|---|---|---:|---|---:|---:|"]
    for row in results["groups"]["main"]:
        counts = row["counts"]
        lines.append(f"| {row['file']} | {row['parse_status']} | {counts.get('chunks',{}).get('count')} | chunks | {counts.get('chunks',{}).get('dim')} | {row['result']} |")
    lines += ["", "## 第 2 组：切片参数 / 空文件", "", f"- 小 chunk KB：`{short(results['groups']['chunking']['small'])}`", f"- 大 chunk KB：`{short(results['groups']['chunking']['large'])}`", f"- 空文件：`{short(results['groups']['chunking']['empty'])}`", f"- 结论：{results['groups']['chunking']['result']}", ""]
    lines += ["## 第 3 组：异常与鲁棒", "", f"- 损坏 PDF：`{short(results['groups']['robust']['broken_pdf'])}`", f"- 超长文档：`{short(results['groups']['robust']['long_doc'])}`", f"- 重复 ingest：`{short(results['groups']['robust']['force_reindex'])}`", f"- 删除 MinIO 源后 force reindex：`{short(results['groups']['robust']['missing_source_reindex'])}`", ""]
    lines += ["## 第 4 组：维度路由", "", f"- `chunk_model_for_dim` 单元脚本：code={results['groups']['dimensions']['unit']['code']}, stdout=`{results['groups']['dimensions']['unit']['stdout'].strip()}`", f"- 非法维度构造：`{short(results['groups']['dimensions']['bad_dim_construct'])}`", "- 1024/3072 端到端：待真实 embedding key 补测。", ""]
    lock = results["groups"]["locking"]
    lines += ["## 第 5 组：Embedding 模型锁定", "", "| 用例 | 状态码/状态 | detail | 结论 |", "|---|---:|---|---:|", f"| 有文档 KB PATCH embedding_model | {lock['patch_locked']['status']} | `{detail(lock['patch_locked'])}` | {'PASS' if lock['patch_locked']['status']==409 else 'FAIL'} |", f"| 空 KB PATCH embedding_model | {lock['patch_empty']['status']} | `{detail(lock['patch_empty'])}` | {'PASS' if lock['patch_empty']['status']==200 else 'FAIL'} |", ""]
    lines += ["## Demo 已验 / 待真实 key 补测", "", "- Demo 已验：txt/md/docx/pdf 解析、1536 mock embedding 入 `chunks`、切片参数、空文件、损坏 PDF、超长文本、force reindex、MinIO 源缺失兜底、模型锁定逻辑、chunk_model_for_dim。", "- 待真实 key 补测：1024/3072 embedding 端到端写入 `chunks_1024` / `chunks_3072` 并检索命中；真实 provider 非法维度返回 `unsupported_embedding_dim` 的 API 构造路径。", ""]
    lines += ["## 缺陷记录", ""]
    if results["defects"]:
        for d in results["defects"]:
            lines.append(f"- {d['severity']} `{d['group']}` {d['name']}: `{short(d['evidence'])}`")
    else:
        lines.append("- 未发现本步阻断性缺陷。")
    lines += ["", "## 结论", ""]
    ok = not results["defects"]
    lines.append(f"- 文档处理与入库链路是否可靠：{'是，demo 主链路通过' if ok else '存在缺陷'}。")
    lines.append(f"- 是否可以进入 Step 4：{'可以' if ok else '不建议'}。")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    main()
